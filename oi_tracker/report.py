"""Build an Excel report (tables + bar charts) from the open-interest history."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .data import contract_sort_key, strikes_around

TITLE = Font(size=14, bold=True)
BOLD = Font(bold=True)
FILLS = {
    "Call": PatternFill("solid", fgColor="C6EFCE"),
    "Put": PatternFill("solid", fgColor="FFC7CE"),
}
CHART_DATES = 5  # how many most recent days each chart compares


def build_report(history: pd.DataFrame, path: Path, strikes_each_side: int = 10) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _write_summary(summary, history)

    latest = history["date"].max()
    contracts = sorted(
        history.loc[history["date"] == latest, "contract"].unique(), key=contract_sort_key,
    )
    for contract in contracts:
        _write_contract_sheet(
            wb.create_sheet(contract), history[history["contract"] == contract],
            strikes_each_side,
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def summary_table(history: pd.DataFrame) -> pd.DataFrame:
    """Per-contract totals on the latest date, with change vs. the previous run."""
    totals = history.pivot_table(
        index=["contract", "date"], columns="option_type",
        values="open_interest", aggfunc="sum", fill_value=0,
    ).reindex(columns=["Call", "Put"], fill_value=0).reset_index()
    totals["total"] = totals["Call"] + totals["Put"]
    totals["change"] = totals.groupby("contract")["total"].diff()

    latest = history["date"].max()
    today = history[history["date"] == latest]
    rows = []
    for contract in sorted(today["contract"].unique(), key=contract_sort_key):
        grp = today[today["contract"] == contract]
        t = totals[(totals["contract"] == contract) & (totals["date"] == latest)].iloc[0]
        calls, puts = grp[grp["option_type"] == "Call"], grp[grp["option_type"] == "Put"]
        rows.append({
            "Contract": contract,
            "Underlying": grp["underlying_price"].iloc[0],
            "Call OI": int(t["Call"]),
            "Put OI": int(t["Put"]),
            "Put/Call": round(t["Put"] / t["Call"], 2) if t["Call"] else None,
            "Max call OI strike": _max_strike(calls),
            "Max put OI strike": _max_strike(puts),
            "OI change vs prev.": None if pd.isna(t["change"]) else int(t["change"]),
        })
    return pd.DataFrame(rows)


def _max_strike(df: pd.DataFrame):
    if df.empty or df["open_interest"].max() == 0:
        return None
    return float(df.loc[df["open_interest"].idxmax(), "strike"])


def _write_summary(ws, history: pd.DataFrame) -> None:
    latest = history["date"].max()
    ws["A1"] = f"Open interest summary - {latest:%Y-%m-%d}"
    ws["A1"].font = TITLE
    table = summary_table(history)
    for col, name in enumerate(table.columns, start=1):
        ws.cell(row=3, column=col, value=name).font = BOLD
        ws.column_dimensions[get_column_letter(col)].width = max(12, len(name) + 2)
    for r, row in enumerate(table.itertuples(index=False), start=4):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)


def _write_contract_sheet(ws, history: pd.DataFrame, n: int) -> None:
    latest = history["date"].max()
    price = history[history["date"] == latest]["underlying_price"].iloc[0]
    ws["A1"] = f"{ws.title} options open interest"
    ws["A1"].font = TITLE
    ws["A2"] = f"Underlying {price} on {latest:%Y-%m-%d}; {n} strikes each side of the money"

    # Use the strike window around today's price for every row, so columns
    # always refer to the same strike across days.
    strikes = strikes_around(history["strike"], price, n)
    row = 4
    anchors = []
    for option_type in ("Call", "Put"):
        mask = (history["option_type"] == option_type) & history["strike"].isin(strikes)
        table = (
            history.loc[mask]
            .pivot_table(index="date", columns="strike", values="open_interest", aggfunc="sum")
            .reindex(columns=strikes).fillna(0).astype(int).sort_index()
        )
        row, chart = _write_oi_table(ws, table, option_type, row)
        anchors.append(chart)
        row += 2

    chart_row = row + 1
    for i, chart in enumerate(anchors):
        ws.add_chart(chart, f"{'A' if i == 0 else 'M'}{chart_row}")
    ws.column_dimensions["A"].width = 12


def _write_oi_table(ws, table: pd.DataFrame, option_type: str, row: int):
    """Write a date x strike table starting at ``row``; return (next_row, chart)."""
    title = ws.cell(row=row, column=1, value=f"{option_type.upper()} open interest by strike")
    title.font = BOLD
    title.fill = FILLS[option_type]
    header = row + 1
    ws.cell(row=header, column=1, value="Date").font = BOLD
    for c, strike in enumerate(table.columns, start=2):
        ws.cell(row=header, column=c, value=strike).font = BOLD
    for r, (day, values) in enumerate(table.iterrows(), start=header + 1):
        ws.cell(row=r, column=1, value=f"{day:%Y-%m-%d}")
        for c, value in enumerate(values, start=2):
            ws.cell(row=r, column=c, value=int(value))
    last = header + len(table)

    chart = BarChart()
    chart.title = f"{option_type} open interest vs. strike"
    chart.y_axis.title = "Open interest"
    chart.x_axis.title = "Strike"
    # openpyxl >= 3.1 marks axes as deleted by default, which hides their labels in Excel.
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    first = max(header + 1, last - CHART_DATES + 1)
    data = Reference(ws, min_col=1, max_col=len(table.columns) + 1, min_row=first, max_row=last)
    chart.add_data(data, from_rows=True, titles_from_data=True)
    chart.set_categories(
        Reference(ws, min_col=2, max_col=len(table.columns) + 1, min_row=header, max_row=header)
    )
    chart.height, chart.width = 10, 22
    return last + 1, chart
