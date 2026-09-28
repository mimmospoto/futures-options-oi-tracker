from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from oi_tracker.data import append_history, contract_sort_key, snapshot, strikes_around
from oi_tracker.report import build_report, summary_table


def fake_chain(oi_offset=0):
    """Raw rows shaped like Barchart's API response (``raw`` section)."""
    rows = []
    for strike in range(200, 360, 10):
        rows.append({"strikePrice": strike, "optionType": "Call", "openInterest": strike + oi_offset})
        rows.append({"strikePrice": strike, "optionType": "Put", "openInterest": None})
    return rows


def test_snapshot_cleans_raw_rows():
    snap = snapshot(fake_chain(), "KCZ26", 278.6, date(2026, 9, 25))
    assert list(snap.columns) == [
        "date", "contract", "underlying_price", "option_type", "strike", "open_interest",
    ]
    assert len(snap) == 32
    assert snap.loc[snap.option_type == "Put", "open_interest"].eq(0).all()  # None -> 0
    assert snap.loc[snap.option_type == "Call", "open_interest"].iloc[0] == 200


def test_snapshot_empty_chain():
    assert snapshot([], "KCZ26", 1.0, date(2026, 9, 25)).empty


def test_strikes_around_picks_n_each_side():
    assert strikes_around([100, 110, 120, 130, 140, 150], 125, 2) == [110, 120, 130, 140]
    # Near the edge of the chain we get what exists.
    assert strikes_around([100, 110, 120], 105, 5) == [100, 110, 120]


def test_contract_sort_key_orders_by_expiry():
    symbols = ["KCH27", "KCZ26", "KCN27", "KCK27"]
    assert sorted(symbols, key=contract_sort_key) == ["KCZ26", "KCH27", "KCK27", "KCN27"]


def test_append_history_accumulates_days_and_replaces_reruns(tmp_path):
    path = tmp_path / "history.csv"
    append_history(path, snapshot(fake_chain(), "KCZ26", 278.6, date(2026, 9, 24)))
    append_history(path, snapshot(fake_chain(), "KCZ26", 280.0, date(2026, 9, 25)))
    # Rerunning on the same day overwrites instead of duplicating.
    history = append_history(path, snapshot(fake_chain(5), "KCZ26", 280.0, date(2026, 9, 25)))

    assert history["date"].nunique() == 2
    assert len(history) == 64
    day2 = history[(history.date == "2026-09-25") & (history.option_type == "Call")]
    assert day2["open_interest"].iloc[0] == 205


@pytest.mark.filterwarnings("error")
def test_build_report_end_to_end(tmp_path):
    csv, xlsx = tmp_path / "h.csv", tmp_path / "report.xlsx"
    for day, offset in [(24, 0), (25, 10)]:
        snaps = pd.concat([
            snapshot(fake_chain(offset), "KCH27", 271.25, date(2026, 9, day)),
            snapshot(fake_chain(offset), "KCZ26", 278.6, date(2026, 9, day)),
        ])
        history = append_history(csv, snaps)
    build_report(history, xlsx, strikes_each_side=3)

    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Summary", "KCZ26", "KCH27"]  # expiry order

    ws = wb["KCZ26"]
    header = [c.value for c in ws[5]][:7]
    assert header == ["Date", 250, 260, 270, 280, 290, 300]
    assert [ws.cell(row=r, column=1).value for r in (6, 7)] == ["2026-09-24", "2026-09-25"]
    assert len(ws._charts) == 2
    for chart in ws._charts:
        assert chart.x_axis.delete is False and chart.y_axis.delete is False

    summary = summary_table(history).set_index("Contract")
    assert summary.loc["KCZ26", "OI change vs prev."] == 16 * 10
    assert summary.loc["KCZ26", "Max call OI strike"] == 350
